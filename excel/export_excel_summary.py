"""
Excel Master Workbook Exporter.
Compiles all analytics results into a multi-tab executive Excel workbook:
- Sheet 1: Executive KPI Overview
- Sheet 2: Top Revenue Products (Q1)
- Sheet 3: Monthly Sales Trend (Q3)
- Sheet 4: Category Momentum (Q4)
- Sheet 5: RFM Segments Summary (Q2)
- Sheet 6: Urgent Churn Watchlist (Q6)
Applies professional styling, number formatting, and auto-fit column widths using openpyxl.
"""

import sys
from pathlib import Path
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent.parent
PROCESSED_DIR = BASE_DIR / "data" / "processed"
OUTPUT_EXCEL = BASE_DIR / "excel" / "ecommerce_analytics_master.xlsx"


def export_master_workbook():
    print("=== Generating Master Excel Workbook ===")
    
    # Load processed dataframes
    df_products = pd.read_csv(PROCESSED_DIR / "top_products_summary.csv")
    df_monthly = pd.read_csv(PROCESSED_DIR / "monthly_sales_summary.csv")
    df_category = pd.read_csv(PROCESSED_DIR / "category_quarterly_momentum.csv")
    df_rfm_summary = pd.read_csv(PROCESSED_DIR / "rfm_segment_summary.csv")
    df_churn_watch = pd.read_csv(PROCESSED_DIR / "churn_watchlist_actionable.csv")

    # Create Executive Summary KPI DataFrame
    total_order_rev = df_monthly["monthly_revenue"].sum()
    product_sales_rev = df_products["total_revenue"].sum()
    total_orders = df_monthly["total_orders"].sum()
    aov = round(total_order_rev / total_orders, 2)
    total_customers = df_rfm_summary["customer_count"].sum()
    at_risk_count = len(df_churn_watch)
    at_risk_spend = df_churn_watch["total_lifetime_spend"].sum()

    kpi_data = [
        {"Metric Category": "Financial Performance", "Metric Name": "Total Completed Gross Revenue", "Value": total_order_rev, "Format": "Currency"},
        {"Metric Category": "Financial Performance", "Metric Name": "Total Product Item Sales", "Value": product_sales_rev, "Format": "Currency"},
        {"Metric Category": "Financial Performance", "Metric Name": "Total Gross Profit", "Value": df_products["total_gross_profit"].sum(), "Format": "Currency"},
        {"Metric Category": "Financial Performance", "Metric Name": "Storewide Gross Margin %", "Value": round((df_products["total_gross_profit"].sum() / product_sales_rev) * 100, 2), "Format": "Percent"},
        {"Metric Category": "Order Economics", "Metric Name": "Total Completed Orders", "Value": total_orders, "Format": "Integer"},
        {"Metric Category": "Order Economics", "Metric Name": "Average Order Value (AOV)", "Value": aov, "Format": "Currency"},
        {"Metric Category": "Order Economics", "Metric Name": "Total Product Units Sold", "Value": df_products["units_sold"].sum(), "Format": "Integer"},
        {"Metric Category": "Customer Base", "Metric Name": "Total Active Customers Analyzed", "Value": total_customers, "Format": "Integer"},
        {"Metric Category": "Customer Base", "Metric Name": "VIP Champion Customers Count", "Value": df_rfm_summary[df_rfm_summary["rfm_segment"].str.contains("Champion")]["customer_count"].sum(), "Format": "Integer"},
        {"Metric Category": "Risk & Retention", "Metric Name": "Customers Flagged At Churn Risk", "Value": at_risk_count, "Format": "Integer"},
        {"Metric Category": "Risk & Retention", "Metric Name": "Revenue At Immediate Churn Risk", "Value": at_risk_spend, "Format": "Currency"},
        {"Metric Category": "Risk & Retention", "Metric Name": "Customer Base Churn Risk Ratio", "Value": round((at_risk_count / total_customers) * 100, 2), "Format": "Percent"}
    ]
    df_kpi = pd.DataFrame(kpi_data)

    # Write to Excel with openpyxl engine
    with pd.ExcelWriter(OUTPUT_EXCEL, engine="openpyxl") as writer:
        df_kpi.to_excel(writer, sheet_name="Executive Summary", index=False)
        df_products.to_excel(writer, sheet_name="Top Products (Q1)", index=False)
        df_monthly.to_excel(writer, sheet_name="Monthly Trend (Q3)", index=False)
        df_category.to_excel(writer, sheet_name="Category Momentum (Q4)", index=False)
        df_rfm_summary.to_excel(writer, sheet_name="RFM Segments (Q2)", index=False)
        df_churn_watch.head(500).to_excel(writer, sheet_name="Churn Watchlist (Q6)", index=False)

    # Style workbook with openpyxl
    from openpyxl import load_workbook
    wb = load_workbook(OUTPUT_EXCEL)

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    regular_font = Font(name="Segoe UI", size=10)
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        ws.views.sheetView[0].showGridLines = True

        # Header row formatting
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        ws.row_dimensions[1].height = 28

        # Body formatting & auto-fit width
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for row_idx, cell in enumerate(col, start=1):
                if row_idx > 1:
                    cell.font = regular_font
                    cell.border = thin_border
                    # Number formatting heuristics
                    val_str = str(cell.value or "")
                    if isinstance(cell.value, (int, float)):
                        if "price" in str(ws.cell(1, cell.column).value).lower() or "revenue" in str(ws.cell(1, cell.column).value).lower() or "profit" in str(ws.cell(1, cell.column).value).lower() or "spend" in str(ws.cell(1, cell.column).value).lower() or "aov" in str(ws.cell(1, cell.column).value).lower():
                            cell.number_format = "$#,##0.00"
                        elif "pct" in str(ws.cell(1, cell.column).value).lower() or "growth" in str(ws.cell(1, cell.column).value).lower() or "ratio" in str(ws.cell(1, cell.column).value).lower():
                            cell.number_format = '0.00"%"'
                        elif isinstance(cell.value, int) or cell.value.is_integer():
                            cell.number_format = "#,##0"

                val_len = len(str(cell.value or ""))
                if val_len > max_len:
                    max_len = val_len

            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(OUTPUT_EXCEL)
    print(f"[SUCCESS] Master Excel Workbook created at: {OUTPUT_EXCEL}")
    return OUTPUT_EXCEL


if __name__ == "__main__":
    export_master_workbook()
